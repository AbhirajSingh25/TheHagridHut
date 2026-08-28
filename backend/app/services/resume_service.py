from __future__ import annotations


# ============================================================
# ENVIRONMENT
# ============================================================

import os

from dotenv import load_dotenv

load_dotenv(override=True)


# ============================================================
# STANDARD LIBRARY
# ============================================================

import io
import json
import re
import shutil
import subprocess
import tempfile

from collections import Counter
from typing import Any


# ============================================================
# FASTAPI
# ============================================================

from fastapi import (
    APIRouter,
    File,
    HTTPException,
    UploadFile,
)

from pydantic import BaseModel


# ============================================================
# OPTIONAL PDF LIBRARIES
# ============================================================

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None


try:
    import fitz
except ImportError:
    fitz = None


# ============================================================
# OPTIONAL DOCX LIBRARY
# ============================================================

try:
    from docx import Document
except ImportError:
    Document = None


# ============================================================
# OPTIONAL GROQ AI
# ============================================================

try:
    from groq import Groq
except ImportError:
    Groq = None


# ============================================================
# ROUTER
# ============================================================

router = APIRouter(
    prefix="/resume",
    tags=[
        "Resume Intelligence",
    ],
)


# ============================================================
# CONFIGURATION
# ============================================================

MAX_FILE_SIZE = 15 * 1024 * 1024


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
}


GROQ_API_KEY = os.getenv(
    "GROQ_API_KEY",
    "",
).strip()


AI_MODEL = os.getenv(
    "GROQ_MODEL",
    "openai/gpt-oss-20b",
).strip()


if not AI_MODEL:
    AI_MODEL = "openai/gpt-oss-20b"


print("=" * 60)
print("RESUME SERVICE INITIALIZED")
print(
    "GROQ API KEY:",
    "FOUND" if GROQ_API_KEY else "NOT FOUND",
)
print(
    "GROQ MODEL:",
    AI_MODEL,
)
print("=" * 60)


# ============================================================
# REQUEST MODELS
# ============================================================

class JDMatchRequest(
    BaseModel
):
    resume_text: str
    job_description: str


# ============================================================
# COMMON SKILLS
# ============================================================

COMMON_SKILLS = {

    # PROGRAMMING
    "Python",
    "Java",
    "JavaScript",
    "TypeScript",
    "C",
    "C++",
    "C#",
    "Go",
    "Golang",
    "Rust",
    "PHP",
    "Ruby",
    "Kotlin",
    "Swift",
    "R",
    "MATLAB",

    # FRONTEND
    "HTML",
    "CSS",
    "React",
    "React.js",
    "Next.js",
    "Vue",
    "Vue.js",
    "Angular",
    "Svelte",
    "Redux",
    "Tailwind",
    "Tailwind CSS",
    "Bootstrap",

    # BACKEND
    "Node.js",
    "Node",
    "Express",
    "Express.js",
    "NestJS",
    "FastAPI",
    "Django",
    "Flask",
    "Spring",
    "Spring Boot",
    "Hibernate",
    "REST API",
    "REST APIs",
    "GraphQL",
    "gRPC",
    "Microservices",

    # DATABASE
    "SQL",
    "MySQL",
    "PostgreSQL",
    "MongoDB",
    "Redis",
    "SQLite",
    "Oracle",
    "Firebase",
    "Supabase",
    "Snowflake",
    "BigQuery",

    # CLOUD
    "AWS",
    "Amazon Web Services",
    "Azure",
    "Microsoft Azure",
    "GCP",
    "Google Cloud",
    "Google Cloud Platform",

    # DEVOPS
    "Docker",
    "Kubernetes",
    "Jenkins",
    "GitHub Actions",
    "Terraform",
    "CI/CD",
    "Linux",
    "Nginx",

    # DATA
    "Data Analysis",
    "Data Analytics",
    "Data Engineering",
    "Data Science",
    "ETL",
    "ELT",
    "Excel",
    "Power BI",
    "Tableau",
    "Pandas",
    "NumPy",
    "Matplotlib",
    "PySpark",
    "Spark",
    "Hadoop",
    "Databricks",
    "Azure Data Factory",

    # AI / ML
    "Machine Learning",
    "Deep Learning",
    "Artificial Intelligence",
    "Generative AI",
    "LLM",
    "LLMs",
    "NLP",
    "Computer Vision",
    "TensorFlow",
    "PyTorch",
    "Keras",
    "Scikit-learn",
    "OpenCV",
    "Hugging Face",
    "LangChain",
    "LangGraph",
    "FAISS",
    "Vector Database",
    "RAG",

    # CYBERSECURITY
    "Cybersecurity",
    "Cyber Security",
    "Network Security",
    "SIEM",
    "SOC",
    "Threat Intelligence",
    "Vulnerability Assessment",
    "Penetration Testing",
    "Ethical Hacking",
    "OWASP",
    "Splunk",
    "Wireshark",
    "Zscaler",

    # TOOLS
    "Git",
    "GitHub",
    "GitLab",
    "Postman",
    "Jira",
    "Selenium",
    "Playwright",
    "Pytest",
    "JUnit",
    "Maven",
    "Gradle",

    # OTHER
    "Agile",
    "Scrum",
    "Problem Solving",
    "Communication",
    "Leadership",
    "Teamwork",
    "Project Management",
}


# ============================================================
# ROLE SKILL MAP
# ============================================================

ROLE_SKILLS = {

    "software engineer": [
        "Python",
        "Java",
        "SQL",
        "Git",
        "REST API",
    ],

    "software developer": [
        "Python",
        "Java",
        "JavaScript",
        "SQL",
        "Git",
    ],

    "frontend developer": [
        "HTML",
        "CSS",
        "JavaScript",
        "React",
        "TypeScript",
    ],

    "backend developer": [
        "Python",
        "Java",
        "SQL",
        "FastAPI",
        "REST API",
    ],

    "full stack developer": [
        "React",
        "JavaScript",
        "Python",
        "SQL",
        "Node.js",
    ],

    "data analyst": [
        "SQL",
        "Python",
        "Excel",
        "Power BI",
        "Data Analysis",
    ],

    "data engineer": [
        "Python",
        "SQL",
        "ETL",
        "PySpark",
        "Data Engineering",
    ],

    "data scientist": [
        "Python",
        "SQL",
        "Machine Learning",
        "Pandas",
        "Scikit-learn",
    ],

    "machine learning engineer": [
        "Python",
        "Machine Learning",
        "TensorFlow",
        "PyTorch",
        "Scikit-learn",
    ],

    "ai engineer": [
        "Python",
        "Machine Learning",
        "LLM",
        "NLP",
        "RAG",
    ],

    "cybersecurity analyst": [
        "Cybersecurity",
        "Network Security",
        "SIEM",
        "Threat Intelligence",
        "Vulnerability Assessment",
    ],

    "java developer": [
        "Java",
        "Spring Boot",
        "Hibernate",
        "SQL",
        "Git",
    ],
}


# ============================================================
# SCORE HELPERS
# ============================================================

def clamp_score(
    value: Any,
) -> int:

    try:
        score = int(
            round(
                float(value)
            )
        )

    except (
        TypeError,
        ValueError,
    ):
        return 0

    return max(
        0,
        min(
            100,
            score,
        ),
    )


def unique_list(
    values: list[Any],
) -> list[str]:

    seen = set()
    result = []

    for value in values:

        if not isinstance(
            value,
            str,
        ):
            continue

        cleaned = value.strip()

        if not cleaned:
            continue

        key = cleaned.lower()

        if key in seen:
            continue

        seen.add(
            key
        )

        result.append(
            cleaned
        )

    return result


# ============================================================
# TEXT CLEANING
# ============================================================

def clean_text(
    text: str,
) -> str:

    if not text:
        return ""

    text = text.replace(
        "\x00",
        " ",
    )

    text = text.replace(
        "\r",
        "\n",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


# ============================================================
# FILE HELPERS
# ============================================================

def get_extension(
    filename: str,
) -> str:

    return os.path.splitext(
        filename.lower()
    )[1]


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf_text(
    file_bytes: bytes,
) -> str:

    extracted_text = ""

    if PdfReader is not None:

        try:

            reader = PdfReader(
                io.BytesIO(
                    file_bytes
                )
            )

            pages = []

            for page in reader.pages:

                try:

                    page_text = (
                        page.extract_text()
                        or ""
                    )

                    if page_text.strip():
                        pages.append(
                            page_text
                        )

                except Exception:
                    continue

            extracted_text = "\n".join(
                pages
            )

        except Exception:
            extracted_text = ""

    if (
        len(
            extracted_text.strip()
        ) < 100
        and
        fitz is not None
    ):

        try:

            document = fitz.open(
                stream=file_bytes,
                filetype="pdf",
            )

            pages = []

            for page in document:

                try:

                    page_text = page.get_text(
                        "text"
                    )

                    if page_text.strip():
                        pages.append(
                            page_text
                        )

                except Exception:
                    continue

            document.close()

            fallback_text = "\n".join(
                pages
            )

            if (
                len(
                    fallback_text.strip()
                )
                >
                len(
                    extracted_text.strip()
                )
            ):
                extracted_text = fallback_text

        except Exception:
            pass

    return clean_text(
        extracted_text
    )


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_docx_text(
    file_bytes: bytes,
) -> str:

    if Document is None:

        raise HTTPException(
            status_code=500,
            detail=(
                "DOCX support is unavailable. "
                "Install python-docx in the active "
                "virtual environment."
            ),
        )

    try:

        document = Document(
            io.BytesIO(
                file_bytes
            )
        )

        parts = []

        for paragraph in document.paragraphs:

            text = (
                paragraph.text
                or ""
            ).strip()

            if text:
                parts.append(
                    text
                )

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = (
                        cell.text
                        or ""
                    ).strip()

                    if value:
                        cells.append(
                            value
                        )

                if cells:
                    parts.append(
                        " | ".join(
                            cells
                        )
                    )

        return clean_text(
            "\n".join(
                parts
            )
        )

    except HTTPException:
        raise

    except Exception as error:

        raise HTTPException(
            status_code=422,
            detail=(
                "Unable to extract text from DOCX file: "
                f"{str(error)}"
            ),
        )


# ============================================================
# DOC EXTRACTION
# ============================================================

def extract_doc_text(
    file_bytes: bytes,
) -> str:

    antiword_path = shutil.which(
        "antiword"
    )

    libreoffice_path = (
        shutil.which(
            "libreoffice"
        )
        or
        shutil.which(
            "soffice"
        )
    )

    with tempfile.TemporaryDirectory() as temp_dir:

        doc_path = os.path.join(
            temp_dir,
            "resume.doc",
        )

        with open(
            doc_path,
            "wb",
        ) as file:
            file.write(
                file_bytes
            )

        if antiword_path:

            try:

                result = subprocess.run(
                    [
                        antiword_path,
                        doc_path,
                    ],
                    capture_output=True,
                    text=True,
                    timeout=30,
                    check=False,
                )

                text = clean_text(
                    result.stdout
                    or ""
                )

                if text:
                    return text

            except Exception:
                pass

        if libreoffice_path:

            try:

                output_dir = os.path.join(
                    temp_dir,
                    "output",
                )

                os.makedirs(
                    output_dir,
                    exist_ok=True,
                )

                subprocess.run(
                    [
                        libreoffice_path,
                        "--headless",
                        "--convert-to",
                        "txt:Text",
                        "--outdir",
                        output_dir,
                        doc_path,
                    ],
                    capture_output=True,
                    text=True,
                    timeout=60,
                    check=False,
                )

                output_file = os.path.join(
                    output_dir,
                    "resume.txt",
                )

                if os.path.exists(
                    output_file
                ):

                    with open(
                        output_file,
                        "r",
                        encoding="utf-8",
                        errors="ignore",
                    ) as file:

                        text = clean_text(
                            file.read()
                        )

                        if text:
                            return text

            except Exception:
                pass

    raise HTTPException(
        status_code=415,
        detail=(
            "Legacy DOC files require LibreOffice "
            "or antiword for text extraction. "
            "Please upload a PDF or DOCX file."
        ),
    )


# ============================================================
# GENERIC FILE EXTRACTION
# ============================================================

def extract_resume_text(
    file_bytes: bytes,
    extension: str,
) -> str:

    if extension == ".pdf":
        return extract_pdf_text(
            file_bytes
        )

    if extension == ".docx":
        return extract_docx_text(
            file_bytes
        )

    if extension == ".doc":
        return extract_doc_text(
            file_bytes
        )

    raise HTTPException(
        status_code=415,
        detail=(
            "Unsupported resume format."
        ),
    )


# ============================================================
# CONTACT DETECTION
# ============================================================

def detect_email(
    text: str,
) -> str | None:

    match = re.search(
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+"
        r"\.[A-Za-z]{2,}\b",
        text,
    )

    if not match:
        return None

    return match.group(
        0
    )


def detect_phone(
    text: str,
) -> str | None:

    patterns = [
        r"\+?\d[\d\s\-\(\)]{8,}\d",
        r"\b\d{10}\b",
    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            text,
        )

        if match:
            return match.group(
                0
            ).strip()

    return None


def detect_linkedin(
    text: str,
) -> str | None:

    match = re.search(
        r"(https?://)?"
        r"(www\.)?"
        r"linkedin\.com"
        r"/[A-Za-z0-9_\-/%?.=]+",
        text,
        re.IGNORECASE,
    )

    if match:
        return match.group(
            0
        )

    return None


def detect_github(
    text: str,
) -> str | None:

    match = re.search(
        r"(https?://)?"
        r"(www\.)?"
        r"github\.com"
        r"/[A-Za-z0-9_\-/%?.=]+",
        text,
        re.IGNORECASE,
    )

    if match:
        return match.group(
            0
        )

    return None


# ============================================================
# SECTION DETECTION
# ============================================================

SECTION_PATTERNS = {

    "summary": [
        "summary",
        "professional summary",
        "profile",
        "career objective",
        "objective",
        "about me",
    ],

    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "internship",
        "internships",
    ],

    "education": [
        "education",
        "academic background",
        "qualification",
        "qualifications",
    ],

    "skills": [
        "skills",
        "technical skills",
        "core skills",
        "competencies",
        "expertise",
    ],

    "projects": [
        "projects",
        "academic projects",
        "personal projects",
        "key projects",
    ],

    "certifications": [
        "certifications",
        "certificates",
        "licenses",
        "credentials",
    ],

    "achievements": [
        "achievements",
        "awards",
        "accomplishments",
    ],
}


def detect_sections(
    text: str,
) -> dict[str, bool]:

    lower_text = text.lower()
    result = {}

    for section, patterns in SECTION_PATTERNS.items():

        result[
            section
        ] = any(
            pattern in lower_text
            for pattern in patterns
        )

    return result


# ============================================================
# SKILL EXTRACTION
# ============================================================

def normalize_skill(
    skill: str,
) -> str:

    value = skill.lower().strip()

    aliases = {

        "react.js":
            "react",

        "node":
            "node.js",

        "express.js":
            "express",

        "vue.js":
            "vue",

        "amazon web services":
            "aws",

        "microsoft azure":
            "azure",

        "google cloud platform":
            "gcp",

        "cyber security":
            "cybersecurity",

        "rest apis":
            "rest api",

        "llms":
            "llm",

        "tailwind css":
            "tailwind",
    }

    return aliases.get(
        value,
        value,
    )


def extract_skills(
    text: str,
) -> list[str]:

    lower_text = text.lower()

    found = []

    sorted_skills = sorted(
        COMMON_SKILLS,
        key=len,
        reverse=True,
    )

    for skill in sorted_skills:

        normalized = skill.lower()

        pattern = (
            r"(?<![A-Za-z0-9+#.])"
            +
            re.escape(
                normalized
            )
            +
            r"(?![A-Za-z0-9+#.])"
        )

        if re.search(
            pattern,
            lower_text,
            re.IGNORECASE,
        ):
            found.append(
                skill
            )

    unique_skills = []
    seen = set()

    for skill in found:

        normalized = normalize_skill(
            skill
        )

        if normalized in seen:
            continue

        seen.add(
            normalized
        )

        unique_skills.append(
            skill
        )

    return sorted(
        unique_skills,
        key=str.lower,
    )


# ============================================================
# ACTION VERBS
# ============================================================

ACTION_VERBS = {

    "developed",
    "designed",
    "implemented",
    "built",
    "created",
    "engineered",
    "optimized",
    "improved",
    "automated",
    "deployed",
    "integrated",
    "analyzed",
    "led",
    "managed",
    "collaborated",
    "delivered",
    "reduced",
    "increased",
    "achieved",
    "configured",
    "tested",
    "maintained",
    "architected",
}


def count_action_verbs(
    text: str,
) -> int:

    words = re.findall(
        r"[A-Za-z]+",
        text.lower(),
    )

    return sum(
        1
        for word in words
        if word in ACTION_VERBS
    )


# ============================================================
# QUANTIFIABLE ACHIEVEMENTS
# ============================================================

def count_metrics(
    text: str,
) -> int:

    patterns = [

        r"\b\d+%",

        r"\b\d+\+",

        r"\b\d+\s+(?:users|clients|customers|projects)",

        r"\b\d+x\b",

        r"\b\d+\s*(?:ms|seconds|minutes|hours)\b",

        r"\b\d+\s*(?:kb|mb|gb|tb)\b",

    ]

    count = 0

    for pattern in patterns:

        count += len(
            re.findall(
                pattern,
                text,
                re.IGNORECASE,
            )
        )

    return count


# ============================================================
# RESUME QUALITY ANALYSIS
# ============================================================

def calculate_resume_sections_score(
    sections: dict[str, bool],
) -> tuple[int, list[str]]:

    important_sections = [

        "summary",
        "experience",
        "education",
        "skills",
        "projects",
    ]

    present = sum(
        1
        for section in important_sections
        if sections.get(
            section
        )
    )

    missing = [

        section
        for section in important_sections
        if not sections.get(
            section
        )
    ]

    score = (
        present
        /
        len(
            important_sections
        )
        *
        100
    )

    return (
        clamp_score(
            score
        ),
        missing,
    )


def calculate_contact_score(
    email: str | None,
    phone: str | None,
    linkedin: str | None,
    github: str | None,
) -> tuple[int, list[str]]:

    score = 0
    missing = []

    if email:
        score += 35
    else:
        missing.append(
            "professional email"
        )

    if phone:
        score += 25
    else:
        missing.append(
            "phone number"
        )

    if linkedin:
        score += 20
    else:
        missing.append(
            "LinkedIn profile"
        )

    if github:
        score += 20
    else:
        missing.append(
            "GitHub or portfolio link"
        )

    return (
        clamp_score(
            score
        ),
        missing,
    )


def calculate_skills_score(
    skills: list[str],
) -> int:

    count = len(
        skills
    )

    if count >= 20:
        return 100

    if count >= 15:
        return 90

    if count >= 10:
        return 80

    if count >= 7:
        return 70

    if count >= 5:
        return 60

    if count >= 3:
        return 45

    if count >= 1:
        return 25

    return 0


def calculate_action_score(
    action_count: int,
) -> int:

    if action_count >= 15:
        return 100

    if action_count >= 10:
        return 85

    if action_count >= 6:
        return 70

    if action_count >= 3:
        return 50

    return 25


def calculate_metrics_score(
    metrics_count: int,
) -> int:

    if metrics_count >= 8:
        return 100

    if metrics_count >= 5:
        return 85

    if metrics_count >= 3:
        return 70

    if metrics_count >= 1:
        return 50

    return 20


def calculate_length_score(
    text: str,
) -> int:

    words = len(
        re.findall(
            r"\w+",
            text,
        )
    )

    if 350 <= words <= 1200:
        return 100

    if 250 <= words < 350:
        return 80

    if 1200 < words <= 1600:
        return 75

    if 200 <= words < 250:
        return 65

    if 1600 < words <= 2000:
        return 60

    if words < 200:
        return 40

    return 45


# ============================================================
# ATS SCORE
# ============================================================

def calculate_ats_score(
    sections_score: int,
    contact_score: int,
    skills_score: int,
    action_score: int,
    metrics_score: int,
    length_score: int,
) -> tuple[
    int,
    dict[str, int],
]:

    breakdown = {

        "sections":
            clamp_score(
                sections_score
            ),

        "contact":
            clamp_score(
                contact_score
            ),

        "skills":
            clamp_score(
                skills_score
            ),

        "action_verbs":
            clamp_score(
                action_score
            ),

        "quantification":
            clamp_score(
                metrics_score
            ),

        "content_length":
            clamp_score(
                length_score
            ),
    }

    score = (

        breakdown["sections"]
        * 0.25

        +

        breakdown["contact"]
        * 0.15

        +

        breakdown["skills"]
        * 0.25

        +

        breakdown["action_verbs"]
        * 0.15

        +

        breakdown["quantification"]
        * 0.10

        +

        breakdown["content_length"]
        * 0.10
    )

    return (
        clamp_score(
            score
        ),
        breakdown,
    )


# ============================================================
# ROLE DETECTION
# ============================================================

def calculate_role_match(
    skills: list[str],
) -> list[dict[str, Any]]:

    normalized_skills = {

        normalize_skill(
            skill
        )

        for skill in skills
    }

    role_scores = []

    for role, required_skills in ROLE_SKILLS.items():

        matched = []
        missing = []

        for skill in required_skills:

            normalized = normalize_skill(
                skill
            )

            if normalized in normalized_skills:
                matched.append(
                    skill
                )

            else:
                missing.append(
                    skill
                )

        score = (

            len(
                matched
            )

            /

            len(
                required_skills
            )

            *

            100
        )

        role_scores.append({

            "role":
                role.title(),

            "score":
                clamp_score(
                    score
                ),

            "matched_skills":
                matched,

            "missing_skills":
                missing,

        })

    role_scores.sort(
        key=lambda item:
            item["score"],
        reverse=True,
    )

    return role_scores


# ============================================================
# RESUME RECOMMENDATIONS
# ============================================================

def generate_strengths(
    skills: list[str],
    sections: dict[str, bool],
    action_count: int,
    metrics_count: int,
    contact_score: int,
) -> list[str]:

    strengths = []

    if len(skills) >= 8:

        strengths.append(
            "Strong technical skill coverage across multiple relevant technologies."
        )

    elif len(skills) >= 5:

        strengths.append(
            "Good foundation of relevant technical skills."
        )

    if sections.get(
        "experience"
    ):

        strengths.append(
            "Work or internship experience is clearly represented."
        )

    if sections.get(
        "projects"
    ):

        strengths.append(
            "Projects demonstrate practical application of skills."
        )

    if action_count >= 5:

        strengths.append(
            "Experience descriptions use strong action-oriented language."
        )

    if metrics_count >= 3:

        strengths.append(
            "Resume includes measurable achievements and impact."
        )

    if contact_score >= 80:

        strengths.append(
            "Professional contact and online profile information is well covered."
        )

    if not strengths:

        strengths.append(
            "The resume contains a foundation that can be improved with stronger structure and more specific achievements."
        )

    return unique_list(
        strengths
    )


def generate_weaknesses(
    skills: list[str],
    missing_sections: list[str],
    missing_contact: list[str],
    action_count: int,
    metrics_count: int,
) -> list[str]:

    weaknesses = []

    if len(skills) < 5:

        weaknesses.append(
            "The skills section appears limited and could include more relevant tools and technologies."
        )

    if "summary" in missing_sections:

        weaknesses.append(
            "Add a concise professional summary tailored to your target role."
        )

    if "experience" in missing_sections:

        weaknesses.append(
            "Work experience or internship experience is not clearly identifiable."
        )

    if "projects" in missing_sections:

        weaknesses.append(
            "Add relevant projects to demonstrate practical implementation skills."
        )

    if action_count < 4:

        weaknesses.append(
            "Use stronger action verbs when describing experience and projects."
        )

    if metrics_count < 2:

        weaknesses.append(
            "Add measurable results, percentages, scale, or outcomes to demonstrate impact."
        )

    for item in missing_contact[:2]:

        weaknesses.append(
            f"Add a {item} to improve recruiter accessibility."
        )

    return unique_list(
        weaknesses
    )


def generate_improvements(
    ats_score: int,
    missing_sections: list[str],
    missing_contact: list[str],
    skills: list[str],
    action_count: int,
    metrics_count: int,
) -> list[str]:

    improvements = []

    if ats_score < 70:

        improvements.append(
            "Improve overall ATS structure by using clear section headings and role-relevant keywords."
        )

    for section in missing_sections:

        improvements.append(
            f"Add a clear {section.title()} section."
        )

    if len(skills) < 8:

        improvements.append(
            "Expand the skills section with technologies that are genuinely supported by your experience."
        )

    if action_count < 5:

        improvements.append(
            "Start experience bullets with strong action verbs such as Developed, Implemented, Optimized, or Designed."
        )

    if metrics_count < 3:

        improvements.append(
            "Quantify achievements using percentages, performance improvements, project scale, or business impact."
        )

    for item in missing_contact:

        improvements.append(
            f"Add your {item}."
        )

    if not improvements:

        improvements.append(
            "Tailor the resume for each role by prioritizing the most relevant skills and achievements."
        )

    return unique_list(
        improvements
    )[:10]


# ============================================================
# READINESS
# ============================================================

def calculate_readiness(
    ats_score: int,
    skills_score: int,
    sections_score: int,
    role_scores: list[dict[str, Any]],
) -> int:

    best_role_score = 0

    if role_scores:

        best_role_score = max(

            item.get(
                "score",
                0,
            )

            for item in role_scores

        )

    readiness = (

        ats_score
        * 0.45

        +

        skills_score
        * 0.20

        +

        sections_score
        * 0.15

        +

        best_role_score
        * 0.20
    )

    return clamp_score(
        readiness
    )


# ============================================================
# VERDICT
# ============================================================

def generate_verdict(
    ats_score: int,
    readiness: int,
) -> str:

    combined = (

        ats_score
        * 0.60

        +

        readiness
        * 0.40
    )

    combined = clamp_score(
        combined
    )

    if combined >= 85:

        return (
            "Excellent resume foundation with strong ATS and recruiter readiness."
        )

    if combined >= 75:

        return (
            "Strong resume with good market readiness. Minor improvements can make it more competitive."
        )

    if combined >= 60:

        return (
            "Promising resume, but targeted improvements are recommended before applying aggressively."
        )

    if combined >= 45:

        return (
            "The resume has useful content but needs stronger structure, keywords, and measurable impact."
        )

    return (
        "The resume requires significant improvement in structure, skill presentation, and recruiter-focused content."
    )


# ============================================================
# SAFE JSON PARSING
# ============================================================

def parse_ai_json(
    value: str,
) -> dict[str, Any] | None:

    if not value:
        return None

    cleaned = value.strip()

    cleaned = re.sub(
        r"^```json\s*",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )

    cleaned = re.sub(
        r"^```\s*",
        "",
        cleaned,
    )

    cleaned = re.sub(
        r"\s*```$",
        "",
        cleaned,
    )

    try:

        result = json.loads(
            cleaned
        )

        if isinstance(
            result,
            dict,
        ):
            return result

    except Exception:
        pass

    match = re.search(
        r"\{.*\}",
        cleaned,
        re.DOTALL,
    )

    if not match:
        return None

    try:

        result = json.loads(
            match.group(
                0
            )
        )

        if isinstance(
            result,
            dict,
        ):
            return result

    except Exception:
        return None

    return None


# ============================================================
# AI RESPONSE VALIDATION
# ============================================================

def validate_ai_review(
    ai_result: dict[str, Any],
    fallback: dict[str, Any],
) -> dict[str, Any]:

    list_fields = [

        "strengths",

        "weaknesses",

        "recommended_roles",

        "missing_skills",

        "improvements",

    ]

    for field in list_fields:

        value = ai_result.get(
            field
        )

        if not isinstance(
            value,
            list,
        ):

            ai_result[
                field
            ] = fallback.get(
                field,
                [],
            )

        else:

            ai_result[
                field
            ] = unique_list(
                value
            )

    string_fields = [

        "summary",

        "verdict",

    ]

    for field in string_fields:

        value = ai_result.get(
            field
        )

        if (
            not isinstance(
                value,
                str,
            )
            or
            not value.strip()
        ):

            ai_result[
                field
            ] = fallback.get(
                field,
                "",
            )

        else:

            ai_result[
                field
            ] = value.strip()

    ai_readiness = clamp_score(
        ai_result.get(
            "readiness"
        )
    )

    if ai_readiness == 0:

        ai_readiness = clamp_score(
            fallback.get(
                "readiness",
                0,
            )
        )

    ai_result[
        "readiness"
    ] = ai_readiness

    ai_result[
        "strengths"
    ] = unique_list(

        ai_result.get(
            "strengths",
            [],
        )

        +

        fallback.get(
            "strengths",
            [],
        )

    )[:10]

    ai_result[
        "weaknesses"
    ] = unique_list(

        ai_result.get(
            "weaknesses",
            [],
        )

        +

        fallback.get(
            "weaknesses",
            [],
        )

    )[:10]

    ai_result[
        "improvements"
    ] = unique_list(

        ai_result.get(
            "improvements",
            [],
        )

        +

        fallback.get(
            "improvements",
            [],
        )

    )[:12]

    ai_result[
        "missing_skills"
    ] = unique_list(

        ai_result.get(
            "missing_skills",
            [],
        )

        +

        fallback.get(
            "missing_skills",
            [],
        )

    )[:20]

    ai_result[
        "recommended_roles"
    ] = unique_list(

        ai_result.get(
            "recommended_roles",
            []
        )

        +

        fallback.get(
            "recommended_roles",
            [],
        )

    )[:8]

    return ai_result


# ============================================================
# OPTIONAL AI REVIEW
# ============================================================

def get_ai_review(
    resume_text: str,
    fallback: dict[str, Any],
) -> dict[str, Any]:

    if not GROQ_API_KEY:

        print(
            "Resume AI disabled: GROQ_API_KEY not found."
        )

        return fallback

    if Groq is None:

        print(
            "Resume AI disabled: groq package is not installed."
        )

        return fallback

    prompt = f"""
You are an expert recruiter, ATS specialist, and career advisor.

Analyze the resume below carefully.

Return ONLY valid JSON.

Use exactly this structure:

{{
    "strengths": [],
    "weaknesses": [],
    "summary": "",
    "verdict": "",
    "readiness": 0,
    "recommended_roles": [],
    "missing_skills": [],
    "improvements": []
}}

Rules:

- Do not invent skills, experience, projects, education, or achievements.
- Only use information visible in the resume.
- Keep strengths specific and realistic.
- Keep weaknesses constructive and useful.
- Recommend realistic job roles based on the resume.
- readiness must be an integer from 0 to 100.
- missing_skills should represent realistic skill gaps.
- improvements must be actionable.
- Return valid JSON only.
- Do not include markdown.
- Do not include explanations outside JSON.

Resume:

{resume_text[:14000]}
"""

    messages = [

        {

            "role":
            "system",

            "content":
            (
                "You are a precise ATS and recruitment "
                "analysis engine. Always return valid JSON only."
            ),

        },

        {

            "role":
            "user",

            "content":
            prompt,

        },

    ]

    try:

        print(
            f"Resume AI review using model: "
            f"{AI_MODEL}"
        )

        client = Groq(
            api_key=GROQ_API_KEY
        )

        completion = client.chat.completions.create(

            model=AI_MODEL,

            temperature=0.2,

            messages=messages,

        )

        raw_content = (

            completion
            .choices[0]
            .message
            .content

            or
            ""
        ).strip()

        print(
            "Resume AI response received successfully."
        )

        ai_result = parse_ai_json(
            raw_content
        )

        if not ai_result:

            print(
                "AI response was not valid JSON. "
                "Using fallback analysis."
            )

            return fallback

        return validate_ai_review(
            ai_result,
            fallback,
        )

    except Exception as error:

        error_message = str(
            error
        )

        print(
            "Resume AI review failed:"
        )

        print(
            error_message
        )

        if (

            "model_not_found"
            in error_message.lower()

            or

            "does not exist"
            in error_message.lower()

            or

            "model not found"
            in error_message.lower()

        ):

            fallback_model = (
                "openai/gpt-oss-20b"
            )

            print(
                f"Trying fallback model: "
                f"{fallback_model}"
            )

            try:

                client = Groq(
                    api_key=GROQ_API_KEY
                )

                completion = client.chat.completions.create(

                    model=fallback_model,

                    temperature=0.2,

                    messages=messages,

                )

                raw_content = (

                    completion
                    .choices[0]
                    .message
                    .content

                    or
                    ""
                ).strip()

                ai_result = parse_ai_json(
                    raw_content
                )

                if ai_result:

                    print(
                        f"Fallback model succeeded: "
                        f"{fallback_model}"
                    )

                    return validate_ai_review(
                        ai_result,
                        fallback,
                    )

            except Exception as fallback_error:

                print(
                    "Fallback Groq model failed:"
                )

                print(
                    str(
                        fallback_error
                    )
                )

        print(
            "Using deterministic resume analysis fallback."
        )

        return fallback


# ============================================================
# KEYWORD EXTRACTION
# ============================================================

STOP_WORDS = {

    "and",
    "the",
    "with",
    "for",
    "from",
    "that",
    "this",
    "will",
    "have",
    "your",
    "our",
    "you",
    "are",
    "their",
    "they",
    "into",
    "using",
    "used",
    "role",
    "work",
    "team",
    "experience",
    "years",
    "year",
    "job",
    "candidate",
    "skills",
    "skill",
    "responsibilities",
    "requirements",
    "preferred",
    "strong",
    "ability",
    "knowledge",
    "development",
}


def extract_keywords(
    text: str,
    limit: int = 50,
) -> list[str]:

    words = re.findall(
        r"[A-Za-z][A-Za-z0-9+#.\-]{1,}",
        text.lower(),
    )

    filtered = [

        word

        for word in words

        if (
            len(
                word
            ) >= 3

            and

            word
            not in STOP_WORDS
        )

    ]

    counter = Counter(
        filtered
    )

    return [

        word

        for word, _

        in counter.most_common(
            limit
        )

    ]


# ============================================================
# JD SKILL EXTRACTION
# ============================================================

def extract_matching_skills(
    resume_text: str,
    job_description: str,
) -> dict[str, Any]:

    resume_skills = extract_skills(
        resume_text
    )

    jd_skills = extract_skills(
        job_description
    )

    normalized_resume = {

        normalize_skill(
            skill
        ):
        skill

        for skill in resume_skills

    }

    matched = []
    missing = []

    for skill in jd_skills:

        normalized = normalize_skill(
            skill
        )

        if normalized in normalized_resume:

            matched.append(
                skill
            )

        else:

            missing.append(
                skill
            )

    return {

        "resume_skills":
            resume_skills,

        "jd_skills":
            jd_skills,

        "matched_skills":
            unique_list(
                matched
            ),

        "missing_skills":
            unique_list(
                missing
            ),

    }


# ============================================================
# JD MATCH SCORE
# ============================================================

def calculate_jd_match_score(
    resume_text: str,
    job_description: str,
    matched_skills: list[str],
    jd_skills: list[str],
) -> int:

    resume_lower = (
        resume_text.lower()
    )

    if jd_skills:

        skill_score = (

            len(
                matched_skills
            )

            /

            len(
                jd_skills
            )

            *

            100
        )

    else:

        skill_score = 0

    jd_keywords = extract_keywords(
        job_description,
        limit=80,
    )

    matched_keywords = 0

    for keyword in jd_keywords:

        if keyword in resume_lower:

            matched_keywords += 1

    keyword_score = 0

    if jd_keywords:

        keyword_score = (

            matched_keywords

            /

            len(
                jd_keywords
            )

            *

            100
        )

    experience_score = 50

    if re.search(

        r"\b(?:intern|developer|engineer|analyst|manager|specialist)\b",

        resume_lower,

    ):

        experience_score = 80

    if jd_skills:

        final_score = (

            skill_score
            * 0.60

            +

            keyword_score
            * 0.30

            +

            experience_score
            * 0.10
        )

    else:

        final_score = (

            keyword_score
            * 0.75

            +

            experience_score
            * 0.25
        )

    return clamp_score(
        final_score
    )


# ============================================================
# KEYWORDS TO ADD
# ============================================================

def get_keywords_to_add(
    job_description: str,
    resume_text: str,
    missing_skills: list[str],
) -> list[str]:

    resume_lower = (
        resume_text.lower()
    )

    jd_keywords = extract_keywords(
        job_description,
        limit=100,
    )

    missing_keywords = [

        keyword

        for keyword in jd_keywords

        if keyword not in resume_lower

    ]

    result = unique_list(

        missing_skills

        +

        missing_keywords
    )

    return result[:20]


# ============================================================
# RESUME ANALYSIS ENDPOINT
# ============================================================

@router.post(
    "/analyze"
)
async def analyze_resume(
    file: UploadFile = File(
        ...
    ),
):

    filename = (
        file.filename
        or ""
    ).strip()

    if not filename:

        raise HTTPException(
            status_code=400,
            detail=(
                "Resume file is required."
            ),
        )

    extension = get_extension(
        filename
    )

    if extension not in SUPPORTED_EXTENSIONS:

        raise HTTPException(
            status_code=415,
            detail=(
                "Unsupported resume format. "
                "Please upload a PDF, DOCX, or DOC file."
            ),
        )

    file_bytes = await file.read()

    if not file_bytes:

        raise HTTPException(
            status_code=400,
            detail=(
                "The uploaded resume is empty."
            ),
        )

    if len(
        file_bytes
    ) > MAX_FILE_SIZE:

        raise HTTPException(
            status_code=413,
            detail=(
                "Resume file is too large. "
                "Maximum supported size is 15 MB."
            ),
        )

    try:

        resume_text = extract_resume_text(

            file_bytes=file_bytes,

            extension=extension,

        )

    except HTTPException:
        raise

    except Exception as error:

        raise HTTPException(
            status_code=422,
            detail=(
                "Unable to read resume content: "
                f"{str(error)}"
            ),
        )

    if len(
        resume_text.strip()
    ) < 30:

        raise HTTPException(
            status_code=422,
            detail=(
                "Very little text could be extracted from the resume. "
                "If this is a scanned PDF, upload a text-based PDF or DOCX file."
            ),
        )

    email = detect_email(
        resume_text
    )

    phone = detect_phone(
        resume_text
    )

    linkedin = detect_linkedin(
        resume_text
    )

    github = detect_github(
        resume_text
    )

    sections = detect_sections(
        resume_text
    )

    (
        sections_score,
        missing_sections,
    ) = calculate_resume_sections_score(
        sections
    )

    (
        contact_score,
        missing_contact,
    ) = calculate_contact_score(

        email=email,

        phone=phone,

        linkedin=linkedin,

        github=github,

    )

    skills = extract_skills(
        resume_text
    )

    skills_score = calculate_skills_score(
        skills
    )

    action_count = count_action_verbs(
        resume_text
    )

    action_score = calculate_action_score(
        action_count
    )

    metrics_count = count_metrics(
        resume_text
    )

    metrics_score = calculate_metrics_score(
        metrics_count
    )

    length_score = calculate_length_score(
        resume_text
    )

    (
        ats_score,
        ats_breakdown,
    ) = calculate_ats_score(

        sections_score=sections_score,

        contact_score=contact_score,

        skills_score=skills_score,

        action_score=action_score,

        metrics_score=metrics_score,

        length_score=length_score,

    )

    role_scores = calculate_role_match(
        skills
    )

    recommended_roles = [

        item[
            "role"
        ]

        for item in role_scores[:5]

        if item.get(
            "score",
            0,
        ) > 0

    ]

    role_missing_skills = []

    if role_scores:

        best_role = role_scores[
            0
        ]

        role_missing_skills = best_role.get(
            "missing_skills",
            [],
        )

    missing_skills = unique_list(
        role_missing_skills
    )[:15]

    strengths = generate_strengths(

        skills=skills,

        sections=sections,

        action_count=action_count,

        metrics_count=metrics_count,

        contact_score=contact_score,

    )

    weaknesses = generate_weaknesses(

        skills=skills,

        missing_sections=missing_sections,

        missing_contact=missing_contact,

        action_count=action_count,

        metrics_count=metrics_count,

    )

    improvements = generate_improvements(

        ats_score=ats_score,

        missing_sections=missing_sections,

        missing_contact=missing_contact,

        skills=skills,

        action_count=action_count,

        metrics_count=metrics_count,

    )

    readiness = calculate_readiness(

        ats_score=ats_score,

        skills_score=skills_score,

        sections_score=sections_score,

        role_scores=role_scores,

    )

    verdict = generate_verdict(

        ats_score=ats_score,

        readiness=readiness,

    )

    fallback_review = {

        "strengths":
            strengths,

        "weaknesses":
            weaknesses,

        "summary":
            (
                "Resume analysis completed using ATS structure, "
                "skills, contact details, action verbs, "
                "quantification, and role alignment."
            ),

        "verdict":
            verdict,

        "readiness":
            readiness,

        "recommended_roles":
            recommended_roles,

        "missing_skills":
            missing_skills,

        "improvements":
            improvements,

    }

    ai_review = get_ai_review(

        resume_text=resume_text,

        fallback=fallback_review,

    )

    return {

        "ats_score":
            ats_score,

        "resume_text":
            resume_text,

        "skills":
            skills,

        "missing_skills":
            unique_list(
                ai_review.get(
                    "missing_skills",
                    missing_skills,
                )
            )[:20],

        "readiness":
            clamp_score(
                ai_review.get(
                    "readiness",
                    readiness,
                )
            ),

        "verdict":
            ai_review.get(
                "verdict",
                verdict,
            ),

        "recommended_roles":
            unique_list(
                ai_review.get(
                    "recommended_roles",
                    recommended_roles,
                )
            )[:8],

        "strengths":
            unique_list(
                ai_review.get(
                    "strengths",
                    strengths,
                )
            )[:10],

        "weaknesses":
            unique_list(
                ai_review.get(
                    "weaknesses",
                    weaknesses,
                )
            )[:10],

        "summary":
            ai_review.get(
                "summary",
                fallback_review[
                    "summary"
                ],
            ),

        "improvements":
            unique_list(
                ai_review.get(
                    "improvements",
                    improvements,
                )
            )[:12],

        "filename":
            filename,

        "file_type":
            extension,

        "contacts": {

            "email":
                email,

            "phone":
                phone,

            "linkedin":
                linkedin,

            "github":
                github,

        },

        "sections":
            sections,

        "missing_sections":
            missing_sections,

        "ats_breakdown":
            ats_breakdown,

        "role_matches":
            role_scores[:8],

        "statistics": {

            "skills_count":
                len(
                    skills
                ),

            "action_verbs":
                action_count,

            "quantified_achievements":
                metrics_count,

            "word_count":
                len(
                    re.findall(
                        r"\w+",
                        resume_text,
                    )
                ),

            "character_count":
                len(
                    resume_text
                ),

        },

    }


# ============================================================
# JOB DESCRIPTION MATCHING
# ============================================================

@router.post(
    "/match-jd"
)
async def match_job_description(
    payload: JDMatchRequest,
):

    resume_text = (
        payload.resume_text
        or ""
    ).strip()

    job_description = (
        payload.job_description
        or ""
    ).strip()

    if len(
        resume_text
    ) < 20:

        raise HTTPException(
            status_code=400,
            detail=(
                "Valid resume text is required."
            ),
        )

    if len(
        job_description
    ) < 20:

        raise HTTPException(
            status_code=400,
            detail=(
                "A valid job description is required."
            ),
        )

    skill_analysis = extract_matching_skills(

        resume_text=resume_text,

        job_description=job_description,

    )

    resume_skills = skill_analysis[
        "resume_skills"
    ]

    jd_skills = skill_analysis[
        "jd_skills"
    ]

    matched_skills = skill_analysis[
        "matched_skills"
    ]

    missing_skills = skill_analysis[
        "missing_skills"
    ]

    match_score = calculate_jd_match_score(

        resume_text=resume_text,

        job_description=job_description,

        matched_skills=matched_skills,

        jd_skills=jd_skills,

    )

    keywords_to_add = get_keywords_to_add(

        job_description=job_description,

        resume_text=resume_text,

        missing_skills=missing_skills,

    )

    skill_match_percentage = 0

    if jd_skills:

        skill_match_percentage = clamp_score(

            len(
                matched_skills
            )

            /

            len(
                jd_skills
            )

            *

            100
        )

    if match_score >= 85:

        match_level = (
            "Excellent Match"
        )

    elif match_score >= 70:

        match_level = (
            "Strong Match"
        )

    elif match_score >= 55:

        match_level = (
            "Moderate Match"
        )

    elif match_score >= 40:

        match_level = (
            "Partial Match"
        )

    else:

        match_level = (
            "Low Match"
        )

    return {

        "match_score":
            match_score,

        "missing_skills":
            missing_skills,

        "keywords_to_add":
            keywords_to_add,

        "match_level":
            match_level,

        "skill_match_percentage":
            skill_match_percentage,

        "resume_skills":
            resume_skills,

        "jd_skills":
            jd_skills,

        "matched_skills":
            matched_skills,

        "matched_skills_count":
            len(
                matched_skills
            ),

        "missing_skills_count":
            len(
                missing_skills
            ),

    }


# ============================================================
# RESUME SERVICE HEALTH
# ============================================================

@router.get(
    "/health"
)
def resume_health():

    return {

        "status":
            "healthy",

        "service":
            "Resume Intelligence",

        "pdf_reader":
            PdfReader
            is not None,

        "pymupdf":
            fitz
            is not None,

        "docx_support":
            Document
            is not None,

        "ai_enabled":
            bool(
                GROQ_API_KEY
            )
            and
            Groq
            is not None,

        "ai_model":
            AI_MODEL,

    }